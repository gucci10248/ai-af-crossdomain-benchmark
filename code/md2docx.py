#!/usr/bin/env python3
"""通用 Markdown → DOCX 转换器（python-docx 纯 Python，无 pandoc 依赖）。
支持：# 标题1-4、**粗体** 行内、普通段落、- 无序/1. 有序列表、> 引用、| 表格 |、--- 分隔线、
链接 [text](url) 降级为 text (url)。用法：python md2docx.py <in.md> <out.docx>
"""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm


def add_runs(p, text):
    """处理 **粗体** 与 [text](url) 行内标记。"""
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not chunk:
            continue
        r = p.add_run(chunk)
        if i % 2 == 1:
            r.bold = True


def convert(md_path, out_path):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    for s in d.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.2)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        ln = lines[i]
        # 表格块
        if ln.strip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            t = d.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for j, htxt in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                r = cell.paragraphs[0].add_run(re.sub(r"\*\*", "", htxt))
                r.bold = True
                r.font.size = Pt(9.5)
            for row in rows:
                cells = t.add_row().cells
                for j in range(len(header)):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "")
                    for r in cells[j].paragraphs[0].runs:
                        r.font.size = Pt(9.5)
            continue
        if not ln.strip():
            i += 1
            continue
        if re.match(r"^---+\s*$", ln.strip()):
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            level = len(m.group(1))
            p = d.add_paragraph()
            r = p.add_run(re.sub(r"\*\*", "", m.group(2)))
            r.bold = True
            r.font.size = Pt({1: 16, 2: 14, 3: 12, 4: 11}[level])
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue
        if ln.strip().startswith(">"):
            p = d.add_paragraph()
            r = p.add_run(re.sub(r"^>\s*", "", ln.strip()))
            r.italic = True
            r.font.size = Pt(9.5)
            i += 1
            continue
        m = re.match(r"^\s*-\s+(.*)$", ln)
        if m:
            p = d.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue
        m = re.match(r"^\s*\d+\.\s+(.*)$", ln)
        if m:
            p = d.add_paragraph(style="List Number")
            add_runs(p, m.group(1))
            i += 1
            continue
        p = d.add_paragraph()
        add_runs(p, ln.strip())
        i += 1

    d.save(out_path)
    print("saved:", out_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
